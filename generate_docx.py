"""Generate approach.docx explaining the Smart Reframe pipeline."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# ── Styles helper ─────────────────────────────────────────────
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0, 51, 102)
    return h

def add_code(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Inches(0.3)
    return p

def add_para(text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p

# ══════════════════════════════════════════════════════════════
#  TITLE
# ══════════════════════════════════════════════════════════════
title = doc.add_heading('Smart Video Reframe: 16:9 to 9:16 Pipeline', level=0)
for run in title.runs:
    run.font.color.rgb = RGBColor(0, 51, 102)

doc.add_paragraph(
    'This document explains the complete approach and code for automatically '
    'converting a standard landscape (16:9) video into a portrait (9:16) video '
    'suitable for Instagram Reels, YouTube Shorts, and TikTok. The system '
    'intelligently tracks subjects (people, faces) and keeps them in frame '
    'throughout the video.'
)

# ══════════════════════════════════════════════════════════════
#  OVERVIEW
# ══════════════════════════════════════════════════════════════
add_heading('Overview of the Approach')
doc.add_paragraph(
    'The pipeline uses a 5-step process. Instead of simply cropping the centre '
    'of each frame, it analyses the video to find where the important subjects '
    'are, then creates a smooth camera path that follows them.'
)

doc.add_paragraph('The five steps are:', style='List Bullet')
steps_overview = [
    'Step 1 - Extract key frames from the video at a lower frame rate.',
    'Step 2 - Detect people, faces, and visually important regions in each frame.',
    'Step 3 - Build a smooth crop trajectory using optimisation and Kalman filtering.',
    'Step 4 - Detect scene cuts, apply easing, and interpolate to full frame rate.',
    'Step 5 - Render the final 9:16 portrait video.',
]
for s in steps_overview:
    doc.add_paragraph(s, style='List Bullet')

# ── Pipeline diagram ─────────────────────────────────────────
add_heading('Pipeline Flow', level=2)
doc.add_paragraph(
    'Input Video (16:9) --> [Frame Extraction] --> [Detection & Saliency] '
    '--> [Trajectory Optimisation] --> [Easing & Interpolation] --> '
    '[Rendering] --> Output Video (9:16)'
)

add_heading('Output Folder Structure', level=2)
add_code(
    'workspace/\n'
    '  frames/             - sampled frames + metadata.json\n'
    '  analysis/           - detections.json, flow.json, saliency/\n'
    '  trajectory/         - trajectory.json\n'
    '  crop_data/          - crop_commands.json\n'
    '  output_9x16.mp4     - FINAL OUTPUT'
)

# ══════════════════════════════════════════════════════════════
#  HOW TO RUN
# ══════════════════════════════════════════════════════════════
add_heading('How to Run the Pipeline')
add_code('python run_pipeline.py your_video.mp4')
doc.add_paragraph('Optional arguments:')
add_code(
    'python run_pipeline.py input.mp4 --sample-fps 10 --output my_output.mp4\n\n'
    'Arguments:\n'
    '  input           : Path to the input 16:9 video file\n'
    '  --workspace     : Output folder (default: workspace)\n'
    '  --sample-fps    : Frames per second to sample (default: 5)\n'
    '  --output        : Output filename (default: output_9x16.mp4)\n'
    '  --method        : "opencv" or "sendcmd" (default: opencv)'
)

# ══════════════════════════════════════════════════════════════
#  MASTER PIPELINE (run_pipeline.py)
# ══════════════════════════════════════════════════════════════
add_heading('Master Pipeline (run_pipeline.py)')
doc.add_paragraph(
    'This is the main entry point that calls all five steps in order. '
    'It accepts command-line arguments (input video path, sample FPS, etc.), '
    'creates the workspace folders, runs each step sequentially, and prints '
    'a summary at the end with total time and number of frames processed.'
)
add_para('Key points:', bold_prefix='')
doc.add_paragraph(
    'It creates subdirectories: frames/, analysis/, trajectory/, crop_data/ inside workspace/.',
    style='List Bullet')
doc.add_paragraph(
    'Each step reads output from the previous step via JSON files, making the pipeline modular.',
    style='List Bullet')
doc.add_paragraph(
    'Times each step and prints a final summary.', style='List Bullet')

# ══════════════════════════════════════════════════════════════
#  STEP 1
# ══════════════════════════════════════════════════════════════
add_heading('Step 1: Frame Extraction (step1_extract_frames.py)')

add_heading('What it does (Simple Explanation)', level=2)
doc.add_paragraph(
    'A video is just a sequence of images (frames) shown very quickly - typically '
    '24 to 30 frames every second. Analysing every single frame would be very slow '
    'and unnecessary, because consecutive frames are almost identical. '
    'So Step 1 picks out only a few frames per second (e.g., 5 per second instead '
    'of 30). These "sampled" frames are saved as JPEG images for the next steps to work on.'
)

add_heading('Approach', level=2)
doc.add_paragraph(
    '1. Open the video using OpenCV (cv2.VideoCapture).\n'
    '2. Read the video properties: resolution, FPS, total frame count, duration.\n'
    '3. Calculate the sampling interval. For example, if the video is 30 FPS and '
    'we want 5 samples per second, we pick every 6th frame (30/5 = 6).\n'
    '4. Loop through all frames; save only the ones at the sampling interval.\n'
    '5. Save a metadata.json file containing video info and which frame numbers were saved.'
)

add_heading('Code Explanation', level=2)
doc.add_paragraph(
    'The extract_frames() function opens the video, reads its properties '
    '(FPS, resolution, frame count), calculates which frames to sample, '
    'then loops through the video saving only the sampled frames as high-quality JPEGs. '
    'It uses a set for O(1) lookup of which frame indices to save. '
    'Finally it writes a metadata.json file that all subsequent steps will read.'
)
add_code(
    'sample_interval = max(1, int(round(original_fps / sample_fps)))\n'
    'sampled_indices = list(range(0, total_frames, sample_interval))\n'
    '# Example: for 30fps video with sample_fps=5 -> every 6th frame'
)

add_heading('Input / Output', level=2)
add_para(' Input: Video file (e.g. your_video.mp4)', bold_prefix='')
add_para(' Output: workspace/frames/frame_000000.jpg, frame_000006.jpg, ... + metadata.json', bold_prefix='')

# ══════════════════════════════════════════════════════════════
#  STEP 2
# ══════════════════════════════════════════════════════════════
add_heading('Step 2: Detection + Saliency + Optical Flow (step2_analyze_frames.py)')

add_heading('What it does (Simple Explanation)', level=2)
doc.add_paragraph(
    'For each sampled frame, this step answers three questions:\n'
    '  (a) Where are the people and faces? (Detection)\n'
    '  (b) What parts of the image naturally attract the eye? (Saliency)\n'
    '  (c) Which direction is everything moving? (Optical Flow)\n\n'
    'The answers are combined into a single "focus point" - the (x, y) coordinate '
    'where the crop window should be centred.'
)

add_heading('Approach', level=2)

add_para('A) Person Detection (YOLOv8)', bold_prefix='')
doc.add_paragraph(
    'YOLOv8-nano is a fast, lightweight object detector. We use it to find '
    'bounding boxes around people in each frame (class 0 = person). '
    'Only detections with confidence > 40% are kept.'
)

add_para('B) Face Detection (Haar Cascade)', bold_prefix='')
doc.add_paragraph(
    'OpenCV\'s built-in Haar cascade classifier detects faces. It works by '
    'scanning the image at multiple scales looking for face-like patterns. '
    'Faces get the highest priority weight (60%) when computing the focus point, '
    'because in most videos the face is what viewers care about most.'
)

add_para('C) Saliency Map (Spectral Residual)', bold_prefix='')
doc.add_paragraph(
    'A saliency map highlights which parts of an image are visually "interesting" '
    '- bright colours, high contrast, or unusual textures. OpenCV\'s Spectral '
    'Residual method converts the image to the frequency domain, finds the '
    '"unexpected" frequencies, and produces a heatmap. The brightest point '
    'on this map is the saliency peak.'
)

add_para('D) Optical Flow (Farneback)', bold_prefix='')
doc.add_paragraph(
    'Optical flow measures how pixels move between two consecutive frames. '
    'The Farneback method computes a motion vector (dx, dy) for every pixel. '
    'We take the median motion to get the dominant movement direction. '
    'This helps the crop "lead" moving subjects instead of lagging behind.'
)

add_para('E) Weighted Focus Point', bold_prefix='')
doc.add_paragraph(
    'All detections are blended into one focus point using weights:\n'
    '  - Faces: 60% weight (highest priority)\n'
    '  - Person bodies: 25% weight\n'
    '  - Saliency peak: 15% weight (fallback)\n'
    'The result is a single (x, y) coordinate per frame.'
)

add_heading('Code Explanation', level=2)
doc.add_paragraph(
    'The analyze_frames() function loads metadata from Step 1, initialises '
    'YOLOv8 and the Haar cascade, then loops over all sampled frames. For each '
    'frame it: (1) computes a saliency map and finds its peak, (2) runs YOLO to '
    'detect persons, (3) runs Haar cascade to detect faces, (4) computes a '
    'weighted average focus point, and (5) computes optical flow vs the previous '
    'frame. Results are saved as detections.json and flow.json.'
)
add_code(
    '# Weighted focus point formula:\n'
    'focus = (face_centres * 0.60 + person_centres * 0.25 + saliency_peak * 0.15)'
)

# ══════════════════════════════════════════════════════════════
#  STEP 3
# ══════════════════════════════════════════════════════════════
add_heading('Step 3: Trajectory Optimisation + Kalman Smoothing (step3_trajectory.py)')

add_heading('What it does (Simple Explanation)', level=2)
doc.add_paragraph(
    'The raw focus points from Step 2 can be noisy and jumpy. If we used them '
    'directly, the crop window would shake and jitter. Step 3 smooths these '
    'points into a gentle, stable camera path using two techniques:\n'
    '  (a) Global Optimisation - finds the best overall path.\n'
    '  (b) Kalman Filter - adds physics-based smoothing (like a camera on a gimbal).'
)

add_heading('Approach', level=2)

add_para('A) Crop Window Size Calculation', bold_prefix='')
doc.add_paragraph(
    'For a 9:16 portrait crop from a 1920x1080 video:\n'
    '  crop_width = 1080 * (9/16) = 607 pixels\n'
    '  crop_height = 1080 pixels (full height)\n'
    'The crop centre X must stay within [304, 1616] so the window '
    'never goes outside the frame.'
)

add_para('B) Global Optimisation (L-BFGS-B)', bold_prefix='')
doc.add_paragraph(
    'We define a cost function with two competing goals:\n'
    '  1. Stay close to the raw focus points (so we track subjects)\n'
    '  2. Minimise frame-to-frame jumps (so the camera is smooth)\n'
    'scipy.optimize.minimize finds the best balance across ALL frames at once. '
    'The lambda weights control the trade-off: higher jitter penalty = smoother path.'
)
add_code(
    'Cost = sum( LAMBDA_DIST * (x[i] - raw_x[i])^2\n'
    '          + LAMBDA_JITTER * (x[i] - x[i-1])^2 )'
)

add_para('C) Gaussian Smoothing', bold_prefix='')
doc.add_paragraph(
    'A Gaussian filter (sigma=2) is applied as a first-pass smoother '
    'to remove high-frequency noise from the optimised trajectory.'
)

add_para('D) Kalman Filter', bold_prefix='')
doc.add_paragraph(
    'A 1D constant-velocity Kalman filter treats the crop position as a '
    'physical object with position and velocity. It predicts where the crop '
    'should be next based on its current speed, then corrects using the '
    'actual measurement. This gives a very natural, gimbal-like motion.'
)

add_heading('Code Explanation', level=2)
doc.add_paragraph(
    'build_trajectory() loads the raw focus points and optical flow data. '
    'It first adds a flow-based "lead" bias (anticipating subject movement), '
    'then runs L-BFGS-B optimisation, Gaussian smoothing, and finally the '
    'Kalman filter. The result is a smooth (smooth_x, smooth_y) per frame, '
    'saved as trajectory.json.'
)

# ══════════════════════════════════════════════════════════════
#  STEP 4
# ══════════════════════════════════════════════════════════════
add_heading('Step 4: Scene Cuts + Easing + Interpolation (step4_easing_interpolation.py)')

add_heading('What it does (Simple Explanation)', level=2)
doc.add_paragraph(
    'Step 3 only produced crop positions for the sampled frames (e.g., every '
    '6th frame). Step 4 fills in the gaps for ALL original frames using cubic '
    'spline interpolation. It also detects scene cuts (hard transitions) so the '
    'crop can jump instantly at cuts instead of smoothly sliding. Finally, it '
    'applies "easing" to make camera movements feel cinematic (slow start, '
    'fast middle, slow end).'
)

add_heading('Approach', level=2)

add_para('A) Scene Cut Detection', bold_prefix='')
doc.add_paragraph(
    'Consecutive frames are compared using colour histograms (HSV). If the '
    'histogram correlation drops below 0.45, a scene cut is detected. At scene '
    'cuts, the crop position should jump immediately rather than interpolate.'
)

add_para('B) Easing Functions', bold_prefix='')
doc.add_paragraph(
    'Within each scene (between cuts), an ease-in-out cubic function is applied. '
    'This remaps time so the crop movement starts slowly, speeds up in the '
    'middle, and slows down at the end - just like a real camera operator would move. '
    'A speed limiter caps maximum movement to 6 pixels per frame to prevent jarring jumps.'
)
add_code(
    '# Ease-in-out cubic:\n'
    'if t < 0.5:  result = 4 * t^3\n'
    'else:        result = 1 - (-2*t + 2)^3 / 2'
)

add_para('C) Cubic Spline Interpolation', bold_prefix='')
doc.add_paragraph(
    'A cubic spline is fitted through the keyframe crop positions. This gives a '
    'smooth curve that is evaluated at every single original frame, producing '
    'crop_x, crop_y, crop_w, crop_h for each frame.'
)

add_heading('Code Explanation', level=2)
doc.add_paragraph(
    'build_crop_commands() loads the trajectory and metadata, detects scene cuts '
    'via histogram comparison, applies easing within each scene segment, then '
    'uses scipy CubicSpline to interpolate to full FPS. The centre_to_crop() '
    'function converts each (cx, cy) centre point into a top-left corner with '
    'clamping. The result is crop_commands.json with one entry per original frame.'
)

# ══════════════════════════════════════════════════════════════
#  STEP 5
# ══════════════════════════════════════════════════════════════
add_heading('Step 5: Rendering the Final Video (step5_render.py)')

add_heading('What it does (Simple Explanation)', level=2)
doc.add_paragraph(
    'This step reads the original video frame by frame, applies the crop from '
    'crop_commands.json, resizes each cropped frame to 1080x1920 (Full HD portrait), '
    'and writes the final output video. The original audio is preserved.'
)

add_heading('Approach', level=2)
doc.add_paragraph('Two rendering methods are available:')

add_para('Method A: OpenCV (Default, Recommended)', bold_prefix='')
doc.add_paragraph(
    '1. Open the original video with OpenCV.\n'
    '2. For each frame, look up its crop rectangle from crop_commands.\n'
    '3. Crop the frame using NumPy array slicing: frame[y:y+h, x:x+w].\n'
    '4. Resize to 1080x1920 using Lanczos interpolation (best quality).\n'
    '5. Write to a temporary video file.\n'
    '6. Use FFmpeg to mux (combine) the video with audio from the original.'
)

add_para('Method B: FFmpeg sendcmd (Advanced)', bold_prefix='')
doc.add_paragraph(
    'Generates a text file of timestamped crop commands and passes it to '
    'FFmpeg\'s sendcmd filter. Faster but less reliable across codecs.'
)

add_heading('Code Explanation', level=2)
doc.add_paragraph(
    'render_final_video() loads metadata and crop commands, then calls either '
    'render_with_opencv() or render_with_ffmpeg_sendcmd(). The OpenCV method '
    'reads each frame, applies the variable crop, resizes with INTER_LANCZOS4, '
    'writes to a temp file, then uses subprocess to call FFmpeg for audio muxing. '
    'The temp file is deleted after muxing.'
)

# ══════════════════════════════════════════════════════════════
#  LIBRARIES
# ══════════════════════════════════════════════════════════════
add_heading('Libraries and Dependencies')

libs = [
    ('OpenCV (cv2)', 'Video reading/writing, image processing, face detection, saliency, optical flow.'),
    ('Ultralytics (YOLOv8)', 'State-of-the-art object detection to find people in frames.'),
    ('NumPy', 'Fast array operations for image manipulation and mathematical computations.'),
    ('SciPy', 'L-BFGS-B optimisation (trajectory), cubic spline interpolation, Gaussian filter.'),
    ('FilterPy', 'Kalman filter implementation for smooth trajectory estimation.'),
    ('tqdm', 'Progress bars for long-running loops.'),
    ('FFmpeg (external)', 'Video encoding and audio muxing via subprocess calls.'),
]
table = doc.add_table(rows=1, cols=2, style='Light Grid Accent 1')
table.rows[0].cells[0].text = 'Library'
table.rows[0].cells[1].text = 'Purpose'
for name, purpose in libs:
    row = table.add_row()
    row.cells[0].text = name
    row.cells[1].text = purpose

# ══════════════════════════════════════════════════════════════
#  KEY CONCEPTS
# ══════════════════════════════════════════════════════════════
add_heading('Key Concepts Explained Simply')

add_heading('What is a Saliency Map?', level=2)
doc.add_paragraph(
    'A saliency map is like a "heat map" of an image showing which parts '
    'your eyes would naturally look at first. Bright, colourful, or high-contrast '
    'areas get high saliency scores. The algorithm converts the image to the '
    'frequency domain (like sound frequencies but for images), finds the '
    '"surprising" parts, and highlights them.'
)

add_heading('What is Optical Flow?', level=2)
doc.add_paragraph(
    'Optical flow tracks how pixels move between two consecutive frames. '
    'If a person walks to the right, the pixels where they appear shift rightward. '
    'We use the median movement across all pixels to get the dominant motion direction. '
    'This helps the crop "anticipate" where subjects are heading.'
)

add_heading('What is a Kalman Filter?', level=2)
doc.add_paragraph(
    'A Kalman filter is like a smart predictor. Imagine tracking a car: '
    'you know its position and speed, so you can predict where it will be next. '
    'When you get a noisy GPS reading, the filter blends its prediction with '
    'the measurement. For our crop position, it means the camera moves smoothly '
    'like it is on a physical gimbal, with inertia.'
)

add_heading('What is Cubic Spline Interpolation?', level=2)
doc.add_paragraph(
    'We only compute crop positions for sampled frames (e.g., every 6th frame). '
    'Cubic spline draws a smooth curve through these known points and lets us '
    'read off the crop position for every frame in between. Unlike straight-line '
    'interpolation, cubic splines produce smooth, natural curves without sharp corners.'
)

add_heading('What is Easing?', level=2)
doc.add_paragraph(
    'Easing makes movements feel natural. Instead of the crop moving at a '
    'constant speed (which looks robotic), ease-in-out makes it start slowly, '
    'accelerate, then decelerate - like how a real camera operator would pan.'
)

# ══════════════════════════════════════════════════════════════
#  SUMMARY
# ══════════════════════════════════════════════════════════════
add_heading('Summary')
doc.add_paragraph(
    'This pipeline takes any standard landscape video and intelligently converts '
    'it to portrait format by:\n'
    '  1. Sampling key frames to save processing time.\n'
    '  2. Using AI (YOLO) and computer vision (saliency, optical flow) to find '
    'where important subjects are.\n'
    '  3. Optimising and smoothing the crop path so the camera movement looks '
    'professional and stable.\n'
    '  4. Applying cinematic easing and handling scene cuts gracefully.\n'
    '  5. Rendering the final high-quality portrait video with original audio.\n\n'
    'The result is a broadcast-quality 9:16 video that keeps subjects centred '
    'and visible throughout, with smooth, professional camera movements.'
)

# ── Save ──────────────────────────────────────────────────────
doc.save('approach.docx')
print("approach.docx generated successfully!")
